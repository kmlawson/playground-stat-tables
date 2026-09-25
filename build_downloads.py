#!/usr/bin/env python3
"""Write the per-book downloads into downloads/:

  <slug>-tables.xlsx   one sheet per table, plus a hyperlinked Contents sheet
  <slug>-directory.md  the Who's Who / directory entries (books that have them)
  <slug>-chronology.md the chronology events (books that have them)

Run with:  uv run --with openpyxl build_downloads.py   (then build_site.py, which links to these files)
"""
import os, re, json, glob
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
import build_site as bs

OUT = os.path.join(bs.HERE, "downloads")
WARN = ("These tables were transcribed by the vision model of Opus 5.5. Before using any of these figures, "
        "you must verify specific statistics with the original source which is linked to whenever possible.")
NUM = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$|^-?\d+(\.\d+)?$")
BOLD = Font(bold=True)


def cell_value(s):
    """Plain numbers become numbers (with a display format that keeps the printed decimals); everything else stays text."""
    s = (s or "").strip()
    if NUM.match(s) and not (len(s) > 1 and s[0] == "0" and s[1] != "."):
        dec = len(s.split(".")[1]) if "." in s else 0
        v = float(s.replace(",", "")) if dec else int(s.replace(",", ""))
        fmt = ("#,##0" if "," in s else "0") + ("." + "0" * dec if dec else "")
        return v, fmt
    return s, None


def put(ws, row):
    """Append a row; text starting with '=' is printed text (e.g. "= 3.30 miles"), never a formula."""
    ws.append(row)
    for c in ws[ws.max_row]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


def link(c, sheet):
    c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{sheet}'!A1")
    c.font = Font(color="1F4A2C", underline="single")


def label(t):
    return (f"Table {t['table_no']}. " if t.get("table_no") else "") + (t.get("title") or "")


def xlsx(book, tables):
    wb = Workbook()
    idx = wb.active
    idx.title = "Contents"
    put(idx, [book["title"] + " — tables"])
    idx["A1"].font = Font(bold=True, size=14)
    put(idx, [book["source"]])
    put(idx, ["Warning: " + WARN])
    put(idx, [])
    put(idx, ["Sheet", "Chapter", "Table", "Printed pages", "Source"])
    for c in idx[5]:
        c.font = BOLD
    for t in tables:
        name = t["id"][:31]
        ws = wb.create_sheet(name)
        put(ws, [label(t)])
        ws["A1"].font = Font(bold=True, size=13)
        meta = [t.get("chapter") or "", "printed pp. " + ", ".join(t.get("printed_pages") or [])]
        put(ws, [" · ".join(x for x in meta if x)])
        if t.get("caption_extra"):
            put(ws, [t["caption_extra"]])
        src = [s for s in t.get("scans", []) if s.get("url")]
        if src:
            put(ws, ["Scan: " + ", ".join(s["url"] for s in src)])
        put(ws, ["← Contents"])
        link(ws.cell(ws.max_row, 1), "Contents")
        for p in t.get("parts", []):
            put(ws, [])
            if p.get("label"):
                put(ws, [p["label"]])
                ws.cell(ws.max_row, 1).font = BOLD
            put(ws, p.get("columns") or [])
            for c in ws[ws.max_row]:
                c.font = BOLD
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for r in p.get("rows", []):
                put(ws, [cell_value(x)[0] for x in r])
                for c, x in zip(ws[ws.max_row], r):
                    fmt = cell_value(x)[1]
                    if fmt:
                        c.number_format = fmt
        for head, key in (("Footnotes", "footnotes"), ("Transcriber's notes", "transcriber_notes")):
            if t.get(key):
                put(ws, [])
                put(ws, [head])
                ws.cell(ws.max_row, 1).font = BOLD
                for n in t[key]:
                    put(ws, [n])
        ws.column_dimensions["A"].width = 34
        for col in "BCDEFGHIJKLMNOPQRSTUVWXYZ":
            ws.column_dimensions[col].width = 14
        put(idx, [name, t.get("chapter") or "", label(t), ", ".join(t.get("printed_pages") or []),
                    src[0]["url"] if src else ""])
        link(idx.cell(idx.max_row, 1), name)
    for col, w in zip("ABCDE", (12, 30, 60, 14, 60)):
        idx.column_dimensions[col].width = w
    path = os.path.join(OUT, f"{book['slug']}-tables.xlsx")
    wb.save(path)
    return path


def md(book, kind):
    p = os.path.join(bs.HERE, "data", f"{book['slug']}-{kind}.json")
    if not os.path.exists(p):
        return None
    entries = json.load(open(p, encoding="utf-8"))
    title = "Who's Who & Directories" if kind == "directory" else "Chronologies"
    lines = [f"# {book['title']} — {title}", "", f"{len(entries):,} entries, LLM-transcribed. "
             + WARN.replace("These tables were", "These entries were").replace("figures", "details").replace("statistics", "details"), ""]
    last_a = last_s = None
    for e in entries:
        if e["a"] != last_a:
            lines += ["", f"## {e['a']}", ""]
            last_a, last_s = e["a"], None
        if e["s"] and e["s"] != last_s:
            lines += [f"### {e['s']}", ""]
            last_s = e["s"]
        src = f" ([{e['l']}]({e['u']}))" if e.get("u") else ""
        lines.append(f"**{e['n']}** {e['t']} — p. {e['p']}{src}")
        lines.append("")
    path = os.path.join(OUT, f"{book['slug']}-{kind}.md")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return path


def main():
    os.makedirs(OUT, exist_ok=True)
    for b in bs.BOOKS:
        tables = bs.load(b)
        for t in tables:
            t["scans"] = bs.scan_links(b, t)
        if tables:
            print(xlsx(b, tables))
        for kind in ("directory", "chronology"):
            r = md(b, kind)
            if r:
                print(r)


if __name__ == "__main__":
    main()
